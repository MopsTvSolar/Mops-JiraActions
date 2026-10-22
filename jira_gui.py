"""
Interface gráfica (Tkinter) para extração de chamados Jira via REST API.

Ao abrir, conecta automaticamente usando URL/e-mail/API Token do arquivo
.env (veja .env.example) — sem tela de login. O ORDER BY da JQL é fixo no
código; "Projeto", "Grupo Solucionador" e "status a considerar" são
seleção múltipla na tela principal.

Uso:
    python jira_gui.py
"""

import logging
import os
import queue
import threading
import tkinter as tk
from datetime import date, datetime
from tkinter import font as tkfont
from tkinter import messagebox, ttk

import calendar
from collections import Counter

import matplotlib
matplotlib.use("TkAgg")
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import matplotlib.pyplot as plt
from openpyxl import load_workbook
import requests

from jira_extractor import (
    JiraExtractorError,
    build_consolidated_report,
    build_daily_report,
    export_report_pdf,
    load_config,
    log,
    run_chamados_a_violar,
    run_chamados_violados,
    run_extracao_completa,
)

FORNECEDOR_STATUS = "Aguardando Fornecedor"

# Partes fixas da JQL (não editáveis na interface).
ORDER_BY_FIXO = "cf[10419] ASC"

# Opções de seleção múltipla exibidas na tela principal.
# Projetos disponíveis (marcado = considerado). "Central de Incidentes" vem
# marcado por padrão, mantendo o comportamento anterior.
PROJETO_OPTIONS = ["Abertura de Chamados", "Central de Incidentes"]
PROJETO_PADRAO_MARCADO = {"Central de Incidentes"}

GRUPO_SOLUCIONADOR_OPTIONS = [
    "CLBR-TI-OPS-OGS-SOLAR-SALESFORCE-N2",
    "CLBR-TI-OPS-OGS SOLAR SALESFORCE",
    "CLBR-TI-OPS-PROD SOLAR SALESFORCE",
]
# Status existentes considerados pela extração (todos vêm marcados por padrão).
STATUS_OPTIONS = [
    "Triagem",
    "Aguardando Suporte",
    "Aguardando Fornecedor",
    "Reaberto",
    "Em atendimento",
    "Aguardando Cliente",
    "Aberto",
    "Encaminhado",
    "Encerrado",
    "Resolvido",
    "Cancelado",
]

# ---------------------------------------------------------------- paleta
BG = "#f1f5f9"
CARD_BG = "#ffffff"
CARD_BORDER = "#e2e8f0"
HEADER_BG = "#1d4ed8"
HEADER_FG = "#ffffff"
TEXT = "#0f172a"
MUTED = "#64748b"
ACCENT = "#2563eb"
ACCENT_ACTIVE = "#1e40af"
ACCENT_DISABLED = "#93b4f5"
SUCCESS = "#16a34a"
DANGER = "#dc2626"
LOG_BG = "#0f172a"
LOG_FG = "#e2e8f0"
LOG_ERROR = "#f87171"
LOG_WARN = "#fbbf24"
LOG_MUTED = "#94a3b8"

# ---------------------------------------------------------------- cache (Jornada)
_JORNADA_CACHE_BASE = os.getenv("CACHE_DIR") or r"C:\Users\edufe\OneDrive\Área de Trabalho\Cache"
_JORNADA_CACHE_DIR = os.path.join(_JORNADA_CACHE_BASE, "Solar")

_JORNADA_COL_GRUPO = "Grupo Solucionador"
_JORNADA_COL_CRIACAO = "Data de Criação"
_JORNADA_COL_CLASSIF = "Classificação"
_JORNADA_COL_SUBCLASS = "Sub-Classificação"
_JORNADA_COL_EPS = "PROP_Site"

_JORNADA_GRUPOS_N1_N2 = [
    "CLBR-TI-OPS-OGS-SOLAR-SALESFORCE-N2",
    "CLBR-TI-OPS-OGS SOLAR SALESFORCE",
]

_MESES_PT = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]


def _jornada_parse_data(valor):
    if not valor:
        return None
    try:
        return datetime.strptime(str(valor)[:19], "%Y-%m-%dT%H:%M:%S")
    except ValueError:
        return None


def _ler_cache_jornada(inicio, fim):
    """Lê planilhas xlsx de _JORNADA_CACHE_DIR, filtra por Data de Criação em [inicio, fim].

    Retorna (total, total_com_class, Counter(classificacao),
             dict{classif: Counter(sub)}, dict{classif: Counter(eps)})
    """
    pasta = _JORNADA_CACHE_DIR
    if not os.path.isdir(pasta):
        raise JiraExtractorError(f'Pasta de cache não encontrada: "{pasta}"')

    cols_need = (
        _JORNADA_COL_GRUPO, _JORNADA_COL_CRIACAO,
        _JORNADA_COL_CLASSIF, _JORNADA_COL_SUBCLASS, _JORNADA_COL_EPS,
    )

    total = 0
    total_com_class = 0
    class_cnt: Counter = Counter()
    sub_cnt: dict = {}
    eps_cnt: dict = {}

    for nome in os.listdir(pasta):
        if not nome.lower().endswith(".xlsx") or nome.startswith("~$"):
            continue
        try:
            wb = load_workbook(os.path.join(pasta, nome), read_only=True, data_only=True)
        except Exception:
            continue
        try:
            ws = wb.active
            linhas = ws.iter_rows(values_only=True)
            header = next(linhas, None)
            if not header:
                continue
            idx = {c: header.index(c) for c in cols_need if c in header}
            if _JORNADA_COL_GRUPO not in idx or _JORNADA_COL_CRIACAO not in idx:
                continue

            def _val(col, linha, _idx=idx):
                i = _idx.get(col)
                return linha[i] if i is not None and i < len(linha) else None

            for linha in linhas:
                if linha is None or all(v is None for v in linha):
                    continue
                if _val(_JORNADA_COL_GRUPO, linha) not in _JORNADA_GRUPOS_N1_N2:
                    continue
                criacao = _jornada_parse_data(_val(_JORNADA_COL_CRIACAO, linha))
                if not criacao or not (inicio <= criacao.date() <= fim):
                    continue
                total += 1
                cl = _val(_JORNADA_COL_CLASSIF, linha)
                if cl:
                    total_com_class += 1
                    class_cnt[cl] += 1
                    sub = _val(_JORNADA_COL_SUBCLASS, linha)
                    eps = _val(_JORNADA_COL_EPS, linha)
                    if cl not in sub_cnt:
                        sub_cnt[cl] = Counter()
                        eps_cnt[cl] = Counter()
                    if sub:
                        sub_cnt[cl][sub] += 1
                    if eps:
                        eps_cnt[cl][eps] += 1
        finally:
            wb.close()

    return total, total_com_class, class_cnt, sub_cnt, eps_cnt


class QueueLogHandler(logging.Handler):
    """Encaminha registros de log para uma fila, consumida pela thread da UI."""

    def __init__(self, log_queue):
        super().__init__()
        self.log_queue = log_queue

    def emit(self, record):
        self.log_queue.put((record.levelname, self.format(record)))


class Card(tk.Frame):
    """Frame com aparência de cartão: fundo branco e borda sutil."""

    def __init__(self, parent, **kwargs):
        super().__init__(
            parent, bg=CARD_BG, highlightbackground=CARD_BORDER, highlightthickness=1, bd=0, **kwargs
        )


class JornadaWindow(tk.Toplevel):
    """Visão Jornada: Pareto de Classificação + drill-down Sub-Classificação / Localidade."""

    def __init__(self, parent):
        super().__init__(parent)
        self.title("Jornada — Classificação de Chamados")
        self.configure(bg=BG)
        self.geometry("1160x780")
        self.minsize(900, 600)
        self.transient(parent)
        self._dados = None
        self._fig = None
        self._ax_sub = None
        self._ax_eps = None
        self._build_ui()

    def _build_ui(self):
        # ---- filtros ----
        filtro_card = Card(self)
        filtro_card.pack(fill="x", padx=14, pady=(12, 0))
        inner = ttk.Frame(filtro_card, style="Card.TFrame", padding=(14, 10))
        inner.pack(fill="x")

        self._modo_var = tk.StringVar(value="mes")
        for texto, valor in [("De / Para", "periodo"), ("Mês", "mes"), ("Ano completo", "ano")]:
            ttk.Radiobutton(
                inner, text=texto, value=valor, variable=self._modo_var,
                command=self._update_controles,
            ).pack(side="left", padx=(0, 8))

        ttk.Separator(inner, orient="vertical").pack(side="left", fill="y", padx=(6, 10))

        # container que alterna entre os sub-frames de filtro
        self._ctrl_frame = ttk.Frame(inner, style="Card.TFrame")
        self._ctrl_frame.pack(side="left")

        # De / Para
        self._f_periodo = ttk.Frame(self._ctrl_frame, style="Card.TFrame")
        ttk.Label(self._f_periodo, text="De:", style="Muted.Card.TLabel").pack(side="left")
        self._e_de = ttk.Entry(self._f_periodo, width=11)
        self._e_de.insert(0, date.today().strftime("%d/%m/%Y"))
        self._e_de.pack(side="left", padx=(3, 10))
        ttk.Label(self._f_periodo, text="Para:", style="Muted.Card.TLabel").pack(side="left")
        self._e_para = ttk.Entry(self._f_periodo, width=11)
        self._e_para.insert(0, date.today().strftime("%d/%m/%Y"))
        self._e_para.pack(side="left", padx=3)

        # Mês + Ano
        self._f_mes = ttk.Frame(self._ctrl_frame, style="Card.TFrame")
        ttk.Label(self._f_mes, text="Mês:", style="Muted.Card.TLabel").pack(side="left")
        self._cb_mes = ttk.Combobox(self._f_mes, values=_MESES_PT, state="readonly", width=12)
        self._cb_mes.set(_MESES_PT[date.today().month - 1])
        self._cb_mes.pack(side="left", padx=(3, 10))
        ttk.Label(self._f_mes, text="Ano:", style="Muted.Card.TLabel").pack(side="left")
        self._e_ano_mes = ttk.Entry(self._f_mes, width=6)
        self._e_ano_mes.insert(0, str(date.today().year))
        self._e_ano_mes.pack(side="left", padx=3)

        # Ano completo
        self._f_ano = ttk.Frame(self._ctrl_frame, style="Card.TFrame")
        ttk.Label(self._f_ano, text="Ano:", style="Muted.Card.TLabel").pack(side="left")
        self._e_ano = ttk.Entry(self._f_ano, width=6)
        self._e_ano.insert(0, str(date.today().year))
        self._e_ano.pack(side="left", padx=3)

        ttk.Separator(inner, orient="vertical").pack(side="left", fill="y", padx=(10, 8))

        self._btn_gerar = ttk.Button(
            inner, text="Gerar", style="Primary.TButton", command=self._gerar,
        )
        self._btn_gerar.pack(side="left", padx=(0, 10))
        self._lbl_total = ttk.Label(inner, text="", style="Muted.Card.TLabel")
        self._lbl_total.pack(side="left")

        self._update_controles()

        # ---- seletor de classificação (drill-down) ----
        dd = ttk.Frame(self, style="TFrame", padding=(14, 6, 14, 0))
        dd.pack(fill="x")
        ttk.Label(dd, text="Classificação:", style="TLabel").pack(side="left")
        self._cb_classif = ttk.Combobox(dd, values=[], state="readonly", width=34)
        self._cb_classif.pack(side="left", padx=6)
        self._cb_classif.bind("<<ComboboxSelected>>", lambda e: self._update_drilldown())

        # ---- área de gráficos ----
        self._chart_frame = ttk.Frame(self, style="TFrame")
        self._chart_frame.pack(fill="both", expand=True, padx=14, pady=(6, 14))

    def _update_controles(self):
        modo = self._modo_var.get()
        self._f_periodo.pack_forget()
        self._f_mes.pack_forget()
        self._f_ano.pack_forget()
        if modo == "periodo":
            self._f_periodo.pack(side="left")
        elif modo == "mes":
            self._f_mes.pack(side="left")
        else:
            self._f_ano.pack(side="left")

    def _get_inicio_fim(self):
        modo = self._modo_var.get()
        if modo == "periodo":
            try:
                inicio = datetime.strptime(self._e_de.get().strip(), "%d/%m/%Y").date()
                fim = datetime.strptime(self._e_para.get().strip(), "%d/%m/%Y").date()
            except ValueError:
                messagebox.showwarning("Data inválida", "Use o formato dd/mm/aaaa.", parent=self)
                return None, None
            if inicio > fim:
                messagebox.showwarning("Período inválido", "Data início > data fim.", parent=self)
                return None, None
            return inicio, fim
        if modo == "mes":
            try:
                ano = int(self._e_ano_mes.get().strip())
            except ValueError:
                messagebox.showwarning("Ano inválido", "Digite um ano numérico.", parent=self)
                return None, None
            mes = self._cb_mes.current() + 1
            _, last = calendar.monthrange(ano, mes)
            return date(ano, mes, 1), date(ano, mes, last)
        # ano completo
        try:
            ano = int(self._e_ano.get().strip())
        except ValueError:
            messagebox.showwarning("Ano inválido", "Digite um ano numérico.", parent=self)
            return None, None
        return date(ano, 1, 1), date(ano, 12, 31)

    def _gerar(self):
        inicio, fim = self._get_inicio_fim()
        if inicio is None:
            return
        self._btn_gerar.config(state="disabled")
        self._lbl_total.config(text="Carregando…")

        def worker():
            try:
                dados = _ler_cache_jornada(inicio, fim)
                self.after(0, lambda: self._exibir(dados))
            except JiraExtractorError as e:
                msg = str(e)
                self.after(0, lambda: messagebox.showerror("Erro", msg, parent=self))
            except Exception as e:
                msg = str(e)
                self.after(0, lambda: messagebox.showerror("Erro inesperado", msg, parent=self))
            finally:
                self.after(0, lambda: self._btn_gerar.config(state="normal"))

        threading.Thread(target=worker, daemon=True).start()

    def _exibir(self, dados):
        total, total_com_class, class_cnt, sub_cnt, eps_cnt = dados
        self._dados = dados
        self._lbl_total.config(
            text=f"{total} chamados  ({total_com_class} com Classificação)"
        )

        classificacoes = [c for c, _ in class_cnt.most_common()]
        self._cb_classif["values"] = classificacoes
        if classificacoes:
            self._cb_classif.set(classificacoes[0])

        for w in self._chart_frame.winfo_children():
            w.destroy()
        if self._fig is not None:
            plt.close(self._fig)

        self._fig = plt.Figure(figsize=(13, 8), facecolor="#f1f5f9", tight_layout=False)
        gs = self._fig.add_gridspec(
            2, 2, height_ratios=[1.15, 1],
            hspace=0.55, wspace=0.42,
            left=0.07, right=0.97, top=0.94, bottom=0.07,
        )
        ax_pareto = self._fig.add_subplot(gs[0, :])
        self._ax_sub = self._fig.add_subplot(gs[1, 0])
        self._ax_eps = self._fig.add_subplot(gs[1, 1])

        self._draw_pareto(ax_pareto, class_cnt, total, total_com_class)

        top_class = classificacoes[0] if classificacoes else None
        if top_class:
            self._draw_drilldown(top_class, sub_cnt, eps_cnt)

        canvas = FigureCanvasTkAgg(self._fig, master=self._chart_frame)
        canvas.draw()
        canvas.get_tk_widget().pack(fill="both", expand=True)

    def _draw_pareto(self, ax, class_cnt, total, total_com_class):
        top = class_cnt.most_common(15)
        ax.set_facecolor("#f8fafc")
        if not top:
            ax.text(
                0.5, 0.5, "Nenhum chamado classificado no período",
                ha="center", va="center", transform=ax.transAxes, fontsize=10, color=MUTED,
            )
            ax.set_title(
                f"{total} Chamados  (0 com Classificação)", fontsize=9.5, color=TEXT, pad=6,
            )
            return

        labels = [c for c, _ in top]
        vals = [n for _, n in top]
        total_classif = total_com_class or sum(vals)
        acumulado, soma = [], 0
        for v in vals:
            soma += v
            acumulado.append(round(soma / total_classif * 100, 1) if total_classif else 0)

        x = list(range(len(labels)))
        bar_colors = ["#f59e0b"] + [ACCENT] * (len(labels) - 1)
        ax.bar(x, vals, color=bar_colors, alpha=0.88, zorder=2)
        ax.grid(axis="y", linestyle="--", alpha=0.4, zorder=1)
        ax.set_xticks(x)
        ax.set_xticklabels(labels, rotation=33, ha="right", fontsize=7.5)
        ax.set_ylabel("Frequência", fontsize=8)
        ax.set_title(
            f"{total} Chamados  ({total_com_class} com Classificação)",
            fontsize=9.5, color=TEXT, pad=6,
        )
        max_v = max(vals) if vals else 1
        for xi, v in zip(x, vals):
            ax.text(xi, v + max_v * 0.012, str(v),
                    ha="center", va="bottom", fontsize=7, fontweight="bold", color=TEXT)

        ax2 = ax.twinx()
        ax2.plot(x, acumulado, color="#ef4444", marker="o", markersize=4.5,
                 linewidth=1.8, zorder=3)
        ax2.set_ylabel("% Acumulado", color="#ef4444", fontsize=8)
        ax2.tick_params(axis="y", labelcolor="#ef4444", labelsize=7)
        ax2.set_ylim(0, 115)
        for xi, pct in zip(x, acumulado):
            ax2.annotate(f"{pct}%", (xi, pct),
                         textcoords="offset points", xytext=(0, 5),
                         ha="center", fontsize=6.5, color="#ef4444")

    def _draw_drilldown(self, classificacao, sub_cnt, eps_cnt):
        TOP_N = 9

        def _draw_h(ax, counter, titulo):
            ax.cla()
            ax.set_facecolor("#f8fafc")
            top = (counter or Counter()).most_common(TOP_N)
            if not top:
                ax.text(0.5, 0.5, "Sem dados", ha="center", va="center",
                        transform=ax.transAxes, fontsize=8, color=MUTED)
                ax.set_title(titulo, fontsize=8, pad=3)
                return
            labels = [str(c)[:40] for c, _ in reversed(top)]
            vals = [n for _, n in reversed(top)]
            total_top = sum(vals)
            y = list(range(len(labels)))
            ax.barh(y, vals, color=ACCENT, alpha=0.82, height=0.65)
            ax.set_yticks(y)
            ax.set_yticklabels(labels, fontsize=6.5)
            ax.grid(axis="x", linestyle="--", alpha=0.4)
            ax.set_title(titulo, fontsize=8, pad=3)
            max_v = max(vals) if vals else 1
            for i, v in enumerate(vals):
                pct = round(v / total_top * 100, 1) if total_top else 0
                ax.text(v + max_v * 0.015, i, f"{v} ({pct}%)",
                        va="center", fontsize=6, color=TEXT)

        _draw_h(self._ax_sub, sub_cnt.get(classificacao, Counter()),
                f"Sub-Classificação — {classificacao}")
        _draw_h(self._ax_eps, eps_cnt.get(classificacao, Counter()),
                f"Localidade — {classificacao}")

    def _update_drilldown(self):
        if not self._dados or self._ax_sub is None:
            return
        _, _, _, sub_cnt, eps_cnt = self._dados
        classificacao = self._cb_classif.get()
        if not classificacao:
            return
        self._draw_drilldown(classificacao, sub_cnt, eps_cnt)
        self._fig.canvas.draw_idle()

    def destroy(self):
        if self._fig is not None:
            plt.close(self._fig)
        super().destroy()


class JiraApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Extração de Chamados Jira")
        self.geometry("880x640")
        self.minsize(720, 520)
        self.configure(bg=BG)

        self.config_data = None
        self.job_running = False
        self.log_queue = queue.Queue()

        self._setup_styles()

        handler = QueueLogHandler(self.log_queue)
        handler.setFormatter(logging.Formatter("%(asctime)s  %(message)s", "%H:%M:%S"))
        log.addHandler(handler)
        log.setLevel(logging.INFO)

        self._build_header()
        self.content = ttk.Frame(self, padding=20)
        self.content.pack(fill="both", expand=True)

        self._build_connecting_view()
        self.after(200, self._poll_log_queue)
        self.after(100, self._connect)

    # ---------------------------------------------------------------- estilo
    def _setup_styles(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        base_font = ("Segoe UI", 10)

        style.configure("TFrame", background=BG)
        style.configure("Card.TFrame", background=CARD_BG)
        style.configure("Header.TFrame", background=HEADER_BG)

        style.configure("TLabel", background=BG, foreground=TEXT, font=base_font)
        style.configure("Card.TLabel", background=CARD_BG, foreground=TEXT, font=base_font)
        style.configure("Muted.Card.TLabel", background=CARD_BG, foreground=MUTED, font=("Segoe UI", 9))
        style.configure("Section.Card.TLabel", background=CARD_BG, foreground=MUTED, font=("Segoe UI", 9, "bold"))
        style.configure("Header.TLabel", background=HEADER_BG, foreground=HEADER_FG, font=("Segoe UI", 17, "bold"))
        style.configure("HeaderSub.TLabel", background=HEADER_BG, foreground="#dbeafe", font=("Segoe UI", 10))

        style.configure("TRadiobutton", background=CARD_BG, foreground=TEXT, font=base_font)
        style.map("TRadiobutton", background=[("active", CARD_BG)])

        style.configure("TCheckbutton", background=CARD_BG, foreground=TEXT, font=base_font)
        style.map("TCheckbutton", background=[("active", CARD_BG)])

        style.configure("Primary.TButton", font=("Segoe UI", 10, "bold"), foreground="white",
                         background=ACCENT, padding=10, borderwidth=0)
        style.map(
            "Primary.TButton",
            background=[("disabled", ACCENT_DISABLED), ("active", ACCENT_ACTIVE)],
            foreground=[("disabled", "#eff6ff")],
        )

        style.configure("Secondary.TButton", font=("Segoe UI", 9), foreground=TEXT,
                         background="#e2e8f0", padding=8, borderwidth=0)
        style.map("Secondary.TButton", background=[("active", "#cbd5e1")])

    # ---------------------------------------------------------------- header
    def _build_header(self):
        header = ttk.Frame(self, style="Header.TFrame")
        header.pack(fill="x")

        header_inner = ttk.Frame(header, style="Header.TFrame", padding=(24, 16))
        header_inner.pack(fill="x")
        ttk.Label(header_inner, text="Extração de Chamados Jira", style="Header.TLabel").pack(anchor="w")
        ttk.Label(
            header_inner, text="Central de Incidentes · monitoramento de SLA", style="HeaderSub.TLabel"
        ).pack(anchor="w")

    # ------------------------------------------------------------- conexão
    def _build_connecting_view(self):
        wrapper = ttk.Frame(self.content)
        wrapper.place(relx=0.5, rely=0.5, anchor="center")
        ttk.Label(wrapper, text="Conectando ao Jira...", font=("Segoe UI", 11)).pack()
        self._connecting_wrapper = wrapper

    def _connect(self):
        try:
            self.config_data = load_config()
        except SystemExit:
            messagebox.showerror(
                "Configuração ausente",
                "Variáveis obrigatórias ausentes no .env (JIRA_URL, JIRA_EMAIL, JIRA_API_TOKEN).\n"
                "Copie .env.example para .env e preencha os valores.",
            )
            self.destroy()
            return

        def worker():
            try:
                resp = requests.get(
                    f"{self.config_data['url']}/rest/api/3/myself",
                    auth=(self.config_data["email"], self.config_data["token"]),
                    headers={"Accept": "application/json"},
                    timeout=15,
                )
            except requests.exceptions.RequestException as e:
                self.after(0, lambda: self._connect_failed(f"Não foi possível conectar: {e}"))
                return

            if resp.status_code == 200:
                display_name = resp.json().get("displayName", self.config_data["email"])
                self.after(0, lambda: self._connect_success(display_name))
            elif resp.status_code == 401:
                self.after(0, lambda: self._connect_failed("E-mail ou API Token inválidos (configurados no .env)."))
            else:
                self.after(0, lambda: self._connect_failed(f"Erro inesperado ({resp.status_code})."))

        threading.Thread(target=worker, daemon=True).start()

    def _connect_failed(self, message):
        messagebox.showerror("Falha na conexão", message)

    def _connect_success(self, display_name):
        self._connecting_wrapper.destroy()
        self._build_main_frame(display_name)

    # ------------------------------------------------------------ main view
    def _build_main_frame(self, display_name):
        canvas = tk.Canvas(self.content, bg=BG, highlightthickness=0)
        scrollbar = ttk.Scrollbar(self.content, orient="vertical", command=canvas.yview)
        body = ttk.Frame(canvas, padding=(0, 0, 12, 0))

        body.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        body_window = canvas.create_window((0, 0), window=body, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(body_window, width=e.width))

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        canvas.bind("<Enter>", lambda e: canvas.bind_all("<MouseWheel>", _on_mousewheel))
        canvas.bind("<Leave>", lambda e: canvas.unbind_all("<MouseWheel>"))

        # --- status card ---
        status_card = Card(body)
        status_card.pack(fill="x", pady=(0, 14))
        status_inner = ttk.Frame(status_card, style="Card.TFrame", padding=(16, 12))
        status_inner.pack(fill="x")

        self.status_dot = tk.Label(status_inner, text="●", fg=SUCCESS, bg=CARD_BG, font=("Segoe UI", 12))
        self.status_dot.pack(side="left")
        self.status_label = tk.Label(
            status_inner, text=" Conectado", fg=SUCCESS, bg=CARD_BG, font=("Segoe UI", 10, "bold")
        )
        self.status_label.pack(side="left")
        self.identity_label = tk.Label(
            status_inner,
            text=f"Conectado como {display_name}  ({self.config_data['email']})",
            fg=MUTED, bg=CARD_BG, font=("Segoe UI", 9),
        )
        self.identity_label.pack(side="right")

        # --- filtros card ---
        filters_card = Card(body)
        filters_card.pack(fill="x", pady=(0, 14))
        filters_inner = ttk.Frame(filters_card, style="Card.TFrame", padding=(16, 14))
        filters_inner.pack(fill="x")

        ttk.Label(filters_inner, text="PROJETO", style="Section.Card.TLabel").pack(anchor="w")
        projeto_frame = ttk.Frame(filters_inner, style="Card.TFrame")
        projeto_frame.pack(anchor="w", pady=(4, 12), fill="x")
        self.projeto_vars = {}
        for projeto in PROJETO_OPTIONS:
            var = tk.BooleanVar(value=projeto in PROJETO_PADRAO_MARCADO)
            ttk.Checkbutton(projeto_frame, text=projeto, variable=var).pack(side="left", padx=(0, 16))
            self.projeto_vars[projeto] = var

        ttk.Label(filters_inner, text="GRUPO SOLUCIONADOR", style="Section.Card.TLabel").pack(anchor="w")
        grupo_frame = ttk.Frame(filters_inner, style="Card.TFrame")
        grupo_frame.pack(anchor="w", pady=(4, 12), fill="x")
        self.grupo_vars = {}
        for grupo in GRUPO_SOLUCIONADOR_OPTIONS:
            var = tk.BooleanVar(value=True)
            ttk.Checkbutton(grupo_frame, text=grupo, variable=var).pack(anchor="w")
            self.grupo_vars[grupo] = var

        ttk.Label(filters_inner, text="STATUS A CONSIDERAR", style="Section.Card.TLabel").pack(anchor="w")
        status_frame = ttk.Frame(filters_inner, style="Card.TFrame")
        status_frame.pack(anchor="w", pady=(4, 0), fill="x")
        self.status_vars = {}
        columns = 4
        for i, status in enumerate(STATUS_OPTIONS):
            var = tk.BooleanVar(value=True)
            ttk.Checkbutton(status_frame, text=status, variable=var).grid(
                row=i // columns, column=i % columns, sticky="w", padx=(0, 14), pady=2
            )
            self.status_vars[status] = var

        # --- options + actions card ---
        actions_card = Card(body)
        actions_card.pack(fill="x", pady=(0, 14))
        actions_inner = ttk.Frame(actions_card, style="Card.TFrame", padding=(16, 14))
        actions_inner.pack(fill="x")

        ttk.Label(actions_inner, text="FORMATO DE SAÍDA", style="Section.Card.TLabel").pack(anchor="w")
        options_frame = ttk.Frame(actions_inner, style="Card.TFrame")
        options_frame.pack(anchor="w", pady=(4, 14))
        self.format_var = tk.StringVar(value="both")
        for label, value in [("CSV", "csv"), ("Excel", "excel"), ("Ambos", "both")]:
            ttk.Radiobutton(options_frame, text=label, value=value, variable=self.format_var).pack(
                side="left", padx=(0, 16)
            )

        ttk.Label(actions_inner, text="AÇÕES", style="Section.Card.TLabel").pack(anchor="w")
        buttons_frame = ttk.Frame(actions_inner, style="Card.TFrame")
        buttons_frame.pack(fill="x", pady=(4, 0))

        self.btn_completa = ttk.Button(
            buttons_frame, text="📋  Extração completa", style="Primary.TButton",
            command=self._run_completa,
        )
        self.btn_completa.pack(side="left", padx=(0, 8), fill="x", expand=True)

        self.btn_hoje = ttk.Button(
            buttons_frame, text="⏰  Chamados a violar hoje", style="Primary.TButton",
            command=self._run_violar_hoje,
        )
        self.btn_hoje.pack(side="left", padx=(0, 8), fill="x", expand=True)

        self.btn_amanha = ttk.Button(
            buttons_frame, text="📅  Chamados a violar amanhã", style="Primary.TButton",
            command=self._run_violar_amanha,
        )
        self.btn_amanha.pack(side="left", padx=(0, 8), fill="x", expand=True)

        self.btn_violados = ttk.Button(
            buttons_frame, text="🔴  Chamados violados", style="Primary.TButton",
            command=self._run_violados,
        )
        self.btn_violados.pack(side="left", fill="x", expand=True)

        ttk.Button(
            actions_inner, text="Abrir pasta de saída", style="Secondary.TButton", command=self._open_output_folder
        ).pack(anchor="e", pady=(12, 0))

        # --- report diário card ---
        report_card = Card(body)
        report_card.pack(fill="x", pady=(0, 14))
        report_inner = ttk.Frame(report_card, style="Card.TFrame", padding=(16, 14))
        report_inner.pack(fill="x")

        report_header = ttk.Frame(report_inner, style="Card.TFrame")
        report_header.pack(fill="x")
        ttk.Label(report_header, text="REPORTS", style="Section.Card.TLabel").pack(side="left")

        self.btn_report = ttk.Button(
            report_header, text="📝  Report Diário", style="Primary.TButton", command=self._run_report,
        )
        self.btn_report.pack(side="left", padx=(12, 8))

        self.btn_consolidado = ttk.Button(
            report_header, text="📊  Relatório Consolidado", style="Primary.TButton",
            command=self._abrir_dialogo_consolidado,
        )
        self.btn_consolidado.pack(side="left", padx=(0, 8))

        self.btn_copiar_report = ttk.Button(
            report_header, text="Copiar", style="Secondary.TButton", command=self._copiar_report, state="disabled",
        )
        self.btn_copiar_report.pack(side="left", padx=(0, 8))

        self.btn_limpar_report = ttk.Button(
            report_header, text="Limpar", style="Secondary.TButton", command=self._limpar_report, state="disabled",
        )
        self.btn_limpar_report.pack(side="left", padx=(0, 8))

        self.btn_pdf_report = ttk.Button(
            report_header, text="📄 Exportar PDF", style="Secondary.TButton",
            command=self._exportar_pdf_report, state="disabled",
        )
        self.btn_pdf_report.pack(side="left")

        report_box_frame = tk.Frame(report_inner, bg=CARD_BG)
        report_box_frame.pack(fill="x", pady=(10, 0))

        report_scroll = tk.Scrollbar(report_box_frame)
        report_scroll.pack(side="right", fill="y")
        self.report_box = tk.Text(
            report_box_frame, height=10, wrap="word", font=("Segoe UI", 10),
            bg="#f8fafc", fg=TEXT, relief="flat", padx=10, pady=10, state="disabled",
            yscrollcommand=report_scroll.set,
        )
        self.report_box.pack(side="left", fill="both", expand=True)
        report_scroll.config(command=self.report_box.yview)

        # --- detalhamento card ---
        detalhe_card = Card(body)
        detalhe_card.pack(fill="x", pady=(0, 14))
        detalhe_inner = ttk.Frame(detalhe_card, style="Card.TFrame", padding=(16, 14))
        detalhe_inner.pack(fill="x")

        detalhe_header = ttk.Frame(detalhe_inner, style="Card.TFrame")
        detalhe_header.pack(fill="x")
        ttk.Label(detalhe_header, text="DETALHAMENTO", style="Section.Card.TLabel").pack(side="left")

        self.btn_jornada = ttk.Button(
            detalhe_header, text="📈  Jornada", style="Primary.TButton",
            command=self._abrir_jornada,
        )
        self.btn_jornada.pack(side="left", padx=(12, 0))

        # --- log card ---
        log_card = Card(body)
        log_card.pack(fill="both", expand=True)
        log_header = ttk.Frame(log_card, style="Card.TFrame", padding=(16, 10, 16, 0))
        log_header.pack(fill="x")
        ttk.Label(log_header, text="LOG DE EXECUÇÃO", style="Section.Card.TLabel").pack(anchor="w")

        log_body = tk.Frame(log_card, bg=CARD_BG, padx=16)
        log_body.pack(fill="both", expand=True, pady=(6, 16))

        mono_font = tkfont.Font(family="Consolas", size=9)
        log_scroll = tk.Scrollbar(log_body)
        log_scroll.pack(side="right", fill="y")
        self.log_box = tk.Text(
            log_body, bg=LOG_BG, fg=LOG_FG, insertbackground=LOG_FG, font=mono_font,
            wrap="word", state="disabled", relief="flat", padx=10, pady=10,
            yscrollcommand=log_scroll.set,
        )
        self.log_box.pack(fill="both", expand=True)
        log_scroll.config(command=self.log_box.yview)

        self.log_box.tag_config("ERROR", foreground=LOG_ERROR)
        self.log_box.tag_config("WARNING", foreground=LOG_WARN)
        self.log_box.tag_config("INFO", foreground=LOG_FG)
        self.log_box.tag_config("MUTED", foreground=LOG_MUTED)

        self.action_buttons = [
            self.btn_completa, self.btn_hoje, self.btn_amanha, self.btn_violados,
            self.btn_report, self.btn_consolidado,
        ]

    def _set_status(self, text, color):
        self.status_dot.config(fg=color)
        self.status_label.config(text=f" {text}", fg=color)

    def _open_output_folder(self):
        output_dir = os.path.abspath("output")
        os.makedirs(output_dir, exist_ok=True)
        os.startfile(output_dir)

    def _set_job_running(self, running):
        self.job_running = running
        state = "disabled" if running else "normal"
        for btn in self.action_buttons:
            btn.config(state=state)
        if running:
            self._set_status("Executando...", ACCENT)
        else:
            self._set_status("Conectado", SUCCESS)

    def _run_job(self, target, *args, on_success=None, **kwargs):
        if self.job_running:
            messagebox.showinfo("Aguarde", "Já existe uma extração em andamento.")
            return

        self._set_job_running(True)

        def worker():
            try:
                result = target(*args, **kwargs)
                if on_success is not None:
                    self.after(0, lambda: on_success(result))
            except JiraExtractorError as e:
                log.error(str(e))
            except Exception as e:
                log.error("Erro inesperado: %s", e)
            finally:
                self.after(0, lambda: self._set_job_running(False))

        threading.Thread(target=worker, daemon=True).start()

    def _build_jql(self):
        projetos = [p for p, var in self.projeto_vars.items() if var.get()]
        if not projetos:
            messagebox.showwarning("Projeto", "Selecione ao menos um projeto.")
            return None

        grupos = [g for g, var in self.grupo_vars.items() if var.get()]
        if not grupos:
            messagebox.showwarning("Grupo Solucionador", "Selecione ao menos um Grupo Solucionador.")
            return None

        status_considerar = [s for s, var in self.status_vars.items() if var.get()]
        if not status_considerar:
            messagebox.showwarning("Status a considerar", "Selecione ao menos um status.")
            return None

        projetos_str = ", ".join(f'"{p}"' for p in projetos)
        grupos_str = ", ".join(f'"{g}"' for g in grupos)
        status_str = ", ".join(f'"{s}"' for s in status_considerar)
        parts = [
            f'"Grupo Solucionador[Group Picker (single group)]" IN ({grupos_str})',
            f"project IN ({projetos_str})",
            f"status IN ({status_str})",
        ]

        return " AND ".join(parts) + f" ORDER BY {ORDER_BY_FIXO}"

    def _incluir_fornecedor(self):
        var = self.status_vars.get(FORNECEDOR_STATUS)
        return bool(var and var.get())

    def _run_completa(self):
        jql = self._build_jql()
        if jql is None:
            return
        config = dict(self.config_data, jql=jql)
        self._run_job(
            run_extracao_completa, config, fmt=self.format_var.get(), incluir_fornecedor=self._incluir_fornecedor()
        )

    def _run_violar_hoje(self):
        jql = self._build_jql()
        if jql is None:
            return
        config = dict(self.config_data, jql=jql)
        self._run_job(
            run_chamados_a_violar, config, days_ahead=0, fmt=self.format_var.get(),
            incluir_fornecedor=self._incluir_fornecedor(),
        )

    def _run_violar_amanha(self):
        jql = self._build_jql()
        if jql is None:
            return
        config = dict(self.config_data, jql=jql)
        self._run_job(
            run_chamados_a_violar, config, days_ahead=1, fmt=self.format_var.get(),
            incluir_fornecedor=self._incluir_fornecedor(),
        )

    def _run_violados(self):
        jql = self._build_jql()
        if jql is None:
            return
        config = dict(self.config_data, jql=jql)
        self._run_job(
            run_chamados_violados, config, fmt=self.format_var.get(), incluir_fornecedor=self._incluir_fornecedor()
        )

    def _run_report(self):
        # O report sempre considera os 3 Grupos Solucionador e os dois
        # projetos combinados (INC + PDST), independente dos checkboxes.
        self._run_job(
            build_daily_report, self.config_data, GRUPO_SOLUCIONADOR_OPTIONS, on_success=self._exibir_report
        )

    def _abrir_jornada(self):
        JornadaWindow(self)

    def _abrir_dialogo_consolidado(self):
        if self.job_running:
            messagebox.showinfo("Aguarde", "Já existe uma extração em andamento.")
            return

        dialog = tk.Toplevel(self)
        dialog.title("Relatório Consolidado")
        dialog.configure(bg=CARD_BG)
        dialog.resizable(False, False)
        dialog.transient(self)
        dialog.grab_set()

        frame = ttk.Frame(dialog, style="Card.TFrame", padding=20)
        frame.pack()

        ttk.Label(frame, text="Selecione o período", style="Card.TLabel", font=("Segoe UI", 12, "bold")).grid(
            row=0, column=0, columnspan=2, pady=(0, 14), sticky="w"
        )

        hoje = date.today().strftime("%d/%m/%Y")

        ttk.Label(frame, text="Data início (dd/mm/aaaa):", style="Muted.Card.TLabel").grid(
            row=1, column=0, sticky="w", pady=4
        )
        entry_inicio = ttk.Entry(frame, width=14)
        entry_inicio.insert(0, hoje)
        entry_inicio.grid(row=1, column=1, pady=4, padx=(10, 0))

        ttk.Label(frame, text="Data fim (dd/mm/aaaa):", style="Muted.Card.TLabel").grid(
            row=2, column=0, sticky="w", pady=4
        )
        entry_fim = ttk.Entry(frame, width=14)
        entry_fim.insert(0, hoje)
        entry_fim.grid(row=2, column=1, pady=4, padx=(10, 0))

        def _confirmar():
            try:
                inicio = datetime.strptime(entry_inicio.get().strip(), "%d/%m/%Y").date()
                fim = datetime.strptime(entry_fim.get().strip(), "%d/%m/%Y").date()
            except ValueError:
                messagebox.showwarning("Data inválida", "Use o formato dd/mm/aaaa.", parent=dialog)
                return
            if inicio > fim:
                messagebox.showwarning("Período inválido", "A data início não pode ser depois da data fim.", parent=dialog)
                return
            dialog.destroy()
            self._run_job(
                build_consolidated_report, self.config_data, inicio, fim, on_success=self._exibir_report
            )

        btn_frame = ttk.Frame(frame, style="Card.TFrame")
        btn_frame.grid(row=3, column=0, columnspan=2, pady=(16, 0), sticky="ew")
        ttk.Button(btn_frame, text="Gerar", style="Primary.TButton", command=_confirmar).pack(
            side="left", fill="x", expand=True, padx=(0, 8)
        )
        ttk.Button(btn_frame, text="Cancelar", style="Secondary.TButton", command=dialog.destroy).pack(
            side="left", fill="x", expand=True
        )

        entry_inicio.focus_set()
        entry_inicio.bind("<Return>", lambda e: _confirmar())
        entry_fim.bind("<Return>", lambda e: _confirmar())

    def _exibir_report(self, texto):
        self.report_box.config(state="normal")
        self.report_box.delete("1.0", "end")
        self.report_box.insert("1.0", texto)
        self.report_box.config(state="disabled")
        self.btn_copiar_report.config(state="normal")
        self.btn_limpar_report.config(state="normal")
        self.btn_pdf_report.config(state="normal")

    def _copiar_report(self):
        texto = self.report_box.get("1.0", "end").strip()
        self.clipboard_clear()
        self.clipboard_append(texto)

    def _limpar_report(self):
        self.report_box.config(state="normal")
        self.report_box.delete("1.0", "end")
        self.report_box.config(state="disabled")
        self.btn_copiar_report.config(state="disabled")
        self.btn_limpar_report.config(state="disabled")
        self.btn_pdf_report.config(state="disabled")

    def _exportar_pdf_report(self):
        texto = self.report_box.get("1.0", "end").strip()
        if not texto:
            messagebox.showinfo("Nada para exportar", "Gere um report antes de exportar para PDF.")
            return

        output_dir = os.path.abspath("output")
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        caminho = os.path.join(output_dir, f"relatorio_{timestamp}.pdf")

        try:
            export_report_pdf(texto, caminho, titulo="Relatório")
        except Exception as e:
            log.error("Falha ao exportar PDF: %s", e)
            messagebox.showerror("Erro ao exportar PDF", str(e))
            return

        log.info("PDF salvo em: %s", caminho)

    # -------------------------------------------------------------- log UI
    def _poll_log_queue(self):
        try:
            while True:
                level, message = self.log_queue.get_nowait()
                tag = level if level in ("ERROR", "WARNING") else "INFO"
                self.log_box.config(state="normal")
                self.log_box.insert("end", message + "\n", tag)
                self.log_box.see("end")
                self.log_box.config(state="disabled")
        except queue.Empty:
            pass
        self.after(200, self._poll_log_queue)


if __name__ == "__main__":
    app = JiraApp()
    app.mainloop()
